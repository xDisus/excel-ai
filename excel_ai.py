"""excel_ai: converte uma planilha Excel em texto compacto para uma IA ler.

Uso:   python excel_ai.py planilha.xlsx     (gera planilha.ai.txt ao lado)
       python excel_ai.py planilha.xlsx --celula "Aba!B5"   (so o que essa celula usa)
       python excel_ai.py antiga.xlsx nova.xlsx --diff      (so o que mudou entre as duas)
       python excel_ai.py pasta                             (todas as planilhas da pasta)
       python excel_ai.py --instalar-skill [--forcar]       (instala a skill do Kiro em ~/.kiro/skills)
Teste: python excel_ai.py                   (sem argumento roda o auto-teste)

Roda 100% local. Unica dependencia: openpyxl (pip install openpyxl).
"""
import bisect
import datetime
import io
import os
import re
import shutil
import sys
import zipfile
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

LIMITE_NUMEROS = 10    # sequencias de numeros/datas maiores que isso viram resumo
LIMITE_TEXTOS = 50     # idem para textos (menor que isso costuma ser rotulo; fica inteiro)
AMOSTRA = 3            # quantos primeiros valores o resumo mostra
RATIO_DUPLICADA = 0.8  # fracao de formulas iguais para tratar a aba como copia
ERROS = ("#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")

CELL = r"\$?[A-Z]{1,3}\$?\d+"
REF_OK = re.compile(r"^(?:%s(?::%s)?|\$?[A-Z]{1,3}:\$?[A-Z]{1,3}|\$?\d+:\$?\d+)$" % (CELL, CELL))
PECA_CELULA = re.compile(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$")
PECA_COLUNA = re.compile(r"^(\$?)([A-Z]{1,3})$")
PECA_LINHA = re.compile(r"^(\$?)(\d+)$")
# ponytail: regex pode achar nome de aba dentro de texto entre aspas; raro, so afeta o mapa.
NOME_ABA = r"'((?:[^']|'')+)'|([^\W\d][\w.]*)"  # 'Aba com espaco' ou Aba
RE_ABA = re.compile(r"(?<![#\w.\]])(?:%s)!" % NOME_ABA)
RE_BLOCO = re.compile(r"^([A-Z]+)(\d+)(?::([A-Z]+)(\d+))? (?:\[[^\]]*\] )?(?:\(v [^)]*\) )?(=.*)$")
RE_MATRIZ = re.compile(r"^([A-Z]+\d+) array (\S+) \{(=.*)\}$")
RE_EXCECAO = re.compile(r"^  except ([A-Z]+\d+) (.*)$")
RE_SUMIU = re.compile(r"^([A-Z]+\d+) no formula$")
RE_REF_QUEBRADA = re.compile(r"(?:(?:%s)!)?#REF!" % NOME_ABA)


def tokens(texto):
    # ponytail: ~4 caracteres por token (estimativa). Use tiktoken se precisar de precisao.
    return len(texto) // 4


def ref(r, c):
    return "%s%d" % (get_column_letter(c), r)


def limpa(formula):
    """Tira prefixos internos do Excel (_xlfn.XLOOKUP vira XLOOKUP) e quebras de linha (Alt+Enter)."""
    return re.sub(r"[\r\n]+", " ", re.sub(r"_xlfn\.|_xlws\.|_xlpm\.", "", formula))


def fmt(v):
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, float):
        return "%.10g" % v
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat()
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    return str(v).replace("\r", " ").replace("\n", " ")


def formula_de(v):
    """('f', texto) formula comum, ('a', texto, faixa) matriz, ('d',) tabela de dados, None constante."""
    nome = type(v).__name__  # por nome: DataTableFormula nao existe em openpyxl < 3.1
    if nome == "ArrayFormula":
        texto = v.text if v.text.startswith("=") else "=" + v.text
        return ("a", texto, v.ref)
    if nome == "DataTableFormula":
        return ("d",)
    if isinstance(v, str) and v.startswith("=") and len(v) > 1:
        return ("f", v)
    return None


def peca(p, row, col):
    m = PECA_CELULA.match(p)
    if m:
        cabs, letra, rabs, num = m.groups()
        c, r = column_index_from_string(letra), int(num)
        return ("R%d" % r if rabs else "R[%d]" % (r - row)) + ("C%d" % c if cabs else "C[%d]" % (c - col))
    m = PECA_COLUNA.match(p)
    if m:
        c = column_index_from_string(m.group(2))
        return "C%d" % c if m.group(1) else "C[%d]" % (c - col)
    m = PECA_LINHA.match(p)
    r = int(m.group(2))
    return "R%d" % r if m.group(1) else "R[%d]" % (r - row)


def relativa(formula, row, col):
    """=B2*C2 em E2 vira =R[0]C[-3]*R[0]C[-2]: mesmo texto para a coluna inteira."""
    try:
        itens = Tokenizer(formula).items
    except Exception:
        return formula
    partes = []
    for t in itens:
        v = t.value
        # "[" = referencia estruturada (Tabela1[Col]) ou outro arquivo: fica como esta.
        if t.type == Token.OPERAND and t.subtype == Token.RANGE and "[" not in v:
            aba, sep, alvo = v.rpartition("!")
            if REF_OK.match(alvo):
                v = aba + sep + ":".join(peca(p, row, col) for p in alvo.split(":"))
        partes.append(v)
    return "=" + "".join(partes)


def rotulo(formulas, valores, r, c, dr, dc):
    """Texto de cabecalho mais perto na direcao (dr, dc). Pula ate 2 vazias.
    None = nada achado; "" = achou algo que nao e rotulo (numero, formula, '%')."""
    for _ in range(3):
        r, c = r + dr, c + dc
        if r < 1 or c < 1:
            return None
        if (r, c) in formulas:
            return ""
        v = valores.get((r, c))
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if not isinstance(v, str) or v.startswith("#") or len(re.findall(r"[^\W\d_]", v)) < 2:
            return ""
        return " ".join(re.sub(r"[\[\]]", "", v).split())[:30].strip()
    return None


def valor_curto(v):
    """Valor calculado em poucos caracteres, sem ')' (a dica '(v X)' nao entra na formula)."""
    if v is None:
        return ""
    return " ".join(fmt(v).replace(")", "").split())[:20]


def blocos(formulas, valores, vertical, buraco="blank", contexto=None, cache=None):
    """Junta celulas vizinhas com o mesmo padrao. Uma celula diferente no meio vira 'except'.
    contexto = (todas as formulas, todos os valores) da aba, so para achar rotulos."""
    todas, todos = contexto or (formulas, valores)
    por_linha = defaultdict(list)
    for (r, c) in formulas:
        por_linha[c if vertical else r].append(r if vertical else c)

    saida = []
    for fixo in sorted(por_linha):
        def cel(p):
            return (p, fixo) if vertical else (fixo, p)

        segs = []  # [tipo, chave, inicio, fim, excecoes]

        def junta(tipo, chave, p0, p1):
            if segs and tipo in "fx" and segs[-1][0] == tipo and segs[-1][1] == chave and segs[-1][3] == p0 - 1:
                segs[-1][3] = p1
            else:
                segs.append([tipo, chave, p0, p1, []])

        ant = None
        for p in sorted(por_linha[fixo]):
            if ant is not None and p > ant + 1:
                junta("x", None, ant + 1, p - 1)  # buraco: vazio ou constante
            f = formulas[cel(p)]
            if f[0] == "f":
                junta("f", f[2], p, p)
            else:
                junta(f[0], p, p, p)  # matriz e tabela de dados nunca se juntam
            ant = p

        pilha = []
        for s in segs:
            if s[0] == "f" and len(pilha) >= 2:
                meio, antes = pilha[-1], pilha[-2]
                if antes[0] == "f" and antes[1] == s[1] and meio[2] == meio[3] and meio[0] in "fx":
                    pilha.pop()
                    antes[4].append(meio[2])
                    antes[4].extend(s[4])
                    antes[3] = s[3]
                    continue
            pilha.append(s)

        for tipo, _, p0, p1, excecoes in pilha:
            ini = cel(p0)
            if tipo == "a":
                saida.append(linha_formula(ini, formulas[ini]))
            if tipo != "f":
                continue
            faixa = ref(*ini) if p0 == p1 else ref(*ini) + ":" + ref(*cel(p1))
            if p0 == p1:  # celula sozinha: esquerda; acima so se a esquerda estiver vazia
                nome = rotulo(todas, todos, ini[0], ini[1], 0, -1)
                if nome is None:
                    nome = rotulo(todas, todos, ini[0], ini[1], -1, 0)
            else:
                dr, dc = (-1, 0) if vertical else (0, -1)
                nome = rotulo(todas, todos, ini[0], ini[1], dr, dc)
            if nome:
                faixa += " [%s]" % nome
            dica = valor_curto((cache or {}).get(ini)) if p0 != p1 else ""  # celula sozinha: sem dica (custa token demais)
            if dica:
                faixa += " (v %s)" % dica
            saida.append("%s %s" % (faixa, limpa(formulas[ini][1])))
            for p in excecoes:
                f = formulas.get(cel(p))
                if f and f[0] == "f":
                    saida.append("  except %s %s" % (ref(*cel(p)), limpa(f[1])))
                elif cel(p) in valores:
                    saida.append("  except %s value %s" % (ref(*cel(p)), fmt(valores[cel(p)])))
                else:
                    saida.append("  except %s %s" % (ref(*cel(p)), buraco))
    return saida


def expandir(linhas, base=None):
    """Le as linhas de formulas de volta para {celula: formula}. Prova que nada se perdeu."""
    out = {}
    if linhas and linhas[0].startswith("same formulas as"):
        out = dict(base)
        resto = []
        for linha in linhas[1:]:
            m = RE_SUMIU.match(linha)
            if m:
                out.pop(m.group(1), None)
            else:
                resto.append(linha)
        out.update(expandir(resto))  # 'except ... unchanged' nao mexe na celula da base
        return out
    for linha in linhas:
        m = RE_EXCECAO.match(linha)
        if m:
            if m.group(2).startswith("="):
                out[m.group(1)] = m.group(2)
            else:
                out.pop(m.group(1), None)
            continue
        m = RE_MATRIZ.match(linha)
        if m:
            out[m.group(1)] = "{%s}%s" % (m.group(2), m.group(3))
            continue
        l1, n1, l2, n2, f = RE_BLOCO.match(linha).groups()
        origem = l1 + n1
        c1, r1 = column_index_from_string(l1), int(n1)
        c2, r2 = (column_index_from_string(l2), int(n2)) if l2 else (c1, r1)
        tr = tradutor(f, origem)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                destino = ref(r, c)
                out[destino] = f if destino == origem else tr(destino)
    return out


def tradutor(f, origem):
    """Translator do openpyxl, protegendo ABA!#REF! (o Translator apaga esse pedaco)."""
    guardados = []

    def guarda(m):
        guardados.append(m.group(0))
        return "XLREFERR%dX" % (len(guardados) - 1)

    tr = Translator(RE_REF_QUEBRADA.sub(guarda, f), origin=origem)
    return lambda destino: re.sub(r"XLREFERR(\d+)X", lambda m: guardados[int(m.group(1))],
                                  tr.translate_formula(destino))


def esperado(formulas):
    """{celula: formula como deve sair}; matriz leva a faixa junto, para a conferencia pegar faixa diferente."""
    return {ref(r, c): ("{%s}%s" % (f[2], limpa(f[1])) if f[0] == "a" else limpa(f[1]))
            for (r, c), f in formulas.items() if f[0] in "fa"}


def confere(linhas, meu, base=None):
    """True se as linhas expandem exatamente para as formulas originais."""
    try:
        return expandir(linhas, base) == meu
    except Exception:  # linha que o leitor nao entende = conferencia falhou
        return False


def secao_formulas(formulas, valores, buraco="blank", contexto=None, cache=None):
    if not formulas:
        return []
    return min(blocos(formulas, valores, True, buraco, contexto, cache),
               blocos(formulas, valores, False, buraco, contexto, cache),
               key=lambda b: len("\n".join(b)))


def por_celula(formulas):
    """Plano B: uma linha por formula (sempre sem perda)."""
    return [linha_formula(pos, formulas[pos]) for pos in sorted(formulas) if formulas[pos][0] in "fa"]


def linha_formula(pos, f):
    """Uma celula na secao de formulas, no formato que expandir() le."""
    if f[0] == "a":
        return "%s array %s {%s}" % (ref(*pos), f[2], limpa(f[1]))
    return "%s %s" % (ref(*pos), limpa(f[1]))


def diferenca(formulas, meu, base_nome, base):
    """'same formulas as X' + so as celulas que mudam, tambem agrupadas em blocos."""
    mudou = {pos: f for pos, f in formulas.items()
             if f[0] in "fa" and base.get(ref(*pos)) != limpa(f[1])}
    sumiu = sorted((k for k in base if k not in meu), key=lambda k: (coordinate_from_string(k)[1], k))
    return (["same formulas as '%s', except:" % base_nome]
            + secao_formulas(mudou, {}, buraco="unchanged")
            + ["%s no formula" % k for k in sumiu])


def tipo_valor(v):
    if isinstance(v, bool):
        return "texts"
    if isinstance(v, (int, float)):
        return "numbers"
    if isinstance(v, (datetime.date, datetime.time)):
        return "dates"
    return "texts"


def secao_dados(valores):
    """Rotulos ficam inteiros. Sequencias longas de numeros/datas/textos viram resumo."""
    def itens(vertical):
        out = []
        por_linha = defaultdict(list)
        for (r, c) in valores:
            por_linha[c if vertical else r].append(r if vertical else c)
        for fixo, pos in por_linha.items():
            pos.sort()
            cels_de = [(p, fixo) if vertical else (fixo, p) for p in pos]
            corridas = [[cels_de[0]]]
            for ant, x in zip(cels_de, cels_de[1:]):
                vizinha = (x[0] - ant[0] + x[1] - ant[1]) == 1
                if vizinha and tipo_valor(valores[x]) == tipo_valor(valores[ant]):
                    corridas[-1].append(x)
                else:
                    corridas.append([x])
            for cels in corridas:
                vals = [valores[x] for x in cels]
                tipo = tipo_valor(vals[0])
                faixa = "%s:%s" % (ref(*cels[0]), ref(*cels[-1]))
                primeiros = ", ".join(fmt(v) for v in vals[:AMOSTRA])
                if tipo != "texts" and len(cels) > LIMITE_NUMEROS:
                    try:
                        menor, maior = fmt(min(vals)), fmt(max(vals))
                    except TypeError:  # data misturada com hora
                        menor = maior = "?"
                    out.append(cels[0] + ("%s [%d %s] min %s max %s first: %s" % (
                        faixa, len(vals), tipo, menor, maior, primeiros), True))
                elif tipo == "texts" and len(cels) > LIMITE_TEXTOS:
                    out.append(cels[0] + ("%s [%d texts, %d distinct] first: %s" % (
                        faixa, len(vals), len(set(map(fmt, vals))), primeiros), True))
                else:
                    i = 0
                    while i < len(cels):  # textos iguais seguidos viram uma linha so
                        j = i
                        while j + 1 < len(cels) and fmt(vals[j + 1]) == fmt(vals[i]):
                            j += 1
                        if j - i >= 2:
                            out.append(cels[i] + ("%s:%s %s (x%d)" % (
                                ref(*cels[i]), ref(*cels[j]), fmt(vals[i]), j - i + 1), True))
                        else:
                            out.extend(x + ("%s %s" % (ref(*x), fmt(valores[x])), False) for x in cels[i:j + 1])
                        i = j + 1
        return out

    escolhido = min(itens(True), itens(False), key=len)
    escolhido.sort()
    linhas, linha_atual, r_atual = [], [], None
    for r, c, texto, resumo in escolhido:
        if resumo or r != r_atual:
            if linha_atual:
                linhas.append(" | ".join(linha_atual))
            linha_atual, r_atual = [], r
        if resumo:
            linhas.append(texto)
            r_atual = None
        else:
            linha_atual.append(texto)
    if linha_atual:
        linhas.append(" | ".join(linha_atual))
    return linhas


def tem_vba(fonte):
    if hasattr(fonte, "seek"):
        fonte.seek(0)
    try:
        with zipfile.ZipFile(fonte) as z:
            return "xl/vbaProject.bin" in z.namelist()
    except zipfile.BadZipFile:
        return False


def nomes_definidos(wb):
    return ["%s = %s" % (d.name, d.attr_text) for d in _nomes(wb)]


def _nomes(wb, globais=False):
    dn = wb.defined_names
    lista = list(dn.values()) if hasattr(dn, "values") else list(getattr(dn, "definedName", []))
    for ws in ([] if globais else wb.worksheets):  # nomes com escopo de aba (openpyxl 3.1+)
        lista.extend(getattr(ws, "defined_names", {}).values())
    return lista


RE_TABELA = re.compile(r"^([^\[\]]*)\[(.*)\]$")


def parte_aba(texto):
    """"'P 1'!A1" -> ("P 1", "A1"); "A1" -> (None, "A1")."""
    aba, sep, resto = texto.rpartition("!")
    if not sep:
        return None, texto
    if len(aba) > 1 and aba[0] == aba[-1] == "'":
        aba = aba[1:-1].replace("''", "'")
    return aba, resto


def acha_alvo(wb, texto):
    """'Aba'!B5 / Aba!B5 / B5 (aba ativa) -> (titulo, linha, coluna). ValueError se nao der."""
    aba, cel = parte_aba(texto.strip())
    m = PECA_CELULA.match(cel.upper())
    if not m or column_index_from_string(m.group(2)) > 16384:
        raise ValueError("celula invalida: %s (use Aba!B5, uma celula so)" % texto)
    nomes = {ws.title.lower(): ws.title for ws in wb.worksheets}
    titulo = wb.active.title if aba is None else nomes.get(aba.lower())
    if titulo is None:
        raise ValueError("aba nao encontrada: %s" % aba)
    return titulo, int(m.group(4)), column_index_from_string(m.group(2))


def faixa(texto):
    """'A1:B9', 'A:A', '2:2', 'B5' -> (r1, c1, r2, c2), None = sem limite. None se nao for faixa."""
    partes = texto.upper().split(":")
    if len(partes) == 1:  # sozinho so vale celula: "TIR" e nome, nao a coluna TIR
        if not PECA_CELULA.match(partes[0]):
            return None
        partes = partes * 2
    if len(partes) != 2:
        return None
    (a, b) = partes
    ma, mb = PECA_CELULA.match(a), PECA_CELULA.match(b)
    if ma and mb:
        r1, r2 = sorted((int(ma.group(4)), int(mb.group(4))))
        c1, c2 = sorted((column_index_from_string(ma.group(2)), column_index_from_string(mb.group(2))))
        return r1, c1, r2, c2
    ma, mb = PECA_COLUNA.match(a), PECA_COLUNA.match(b)
    if ma and mb:
        c1, c2 = sorted((column_index_from_string(ma.group(2)), column_index_from_string(mb.group(2))))
        return None, c1, None, c2
    ma, mb = PECA_LINHA.match(a), PECA_LINHA.match(b)
    if ma and mb:
        r1, r2 = sorted((int(ma.group(2)), int(mb.group(2))))
        return r1, None, r2, None
    return None


def precedentes(wb, dados, alvo):
    """Celulas de que alvo depende, recursivo entre abas.
    dados = {aba: (formulas, valores)}. Retorna ({aba: set(pos)}, avisos, nomes usados)."""
    abas = {t.lower(): t for t in dados}
    indice, matrizes, tabelas = {}, defaultdict(list), {}
    for aba, (formulas, valores) in dados.items():
        cols = defaultdict(list)
        for (r, c) in set(formulas) | set(valores):
            cols[c].append(r)
        indice[aba] = {c: sorted(rs) for c, rs in cols.items()}
        for (r, c), f in formulas.items():
            if f[0] == "a":
                r1, c1, r2, c2 = faixa(f[2])
                matrizes[aba].append((r1, c1, r2, c2, (r, c)))
    for ws in wb.worksheets:
        for t in ws.tables.values():
            r1, c1, r2, c2 = faixa(t.ref)
            cab = 1 if t.headerRowCount is None else t.headerRowCount
            tabelas[t.displayName.lower()] = (ws.title, r1 + cab, c1, r2, c2,
                                              [col.name.lower() for col in t.tableColumns])
    nomes = {(None, d.name.lower()): (d.name, d.attr_text) for d in _nomes(wb, globais=True)}
    for ws in wb.worksheets:  # nome com escopo de aba ganha do global nessa aba
        for d in getattr(ws, "defined_names", {}).values():
            nomes[(ws.title, d.name.lower())] = (d.name, d.attr_text)
    avisos, exemplo, usados = Counter(), {}, set()

    def celulas(aba, r1, c1, r2, c2):
        for c, linhas in indice[aba].items():
            if c1 is not None and not c1 <= c <= c2:
                continue
            ini = 0 if r1 is None else bisect.bisect_left(linhas, r1)
            fim = len(linhas) if r2 is None else bisect.bisect_right(linhas, r2)
            for r in linhas[ini:fim]:
                yield (aba, r, c)
        for a1, b1, a2, b2, ancora in matrizes[aba]:
            if (r1 is None or (a1 <= r2 and r1 <= a2)) and (c1 is None or (b1 <= c2 and c1 <= b2)):
                yield (aba, ancora[0], ancora[1])

    def tabela(nome, dentro, aqui, pos):
        if nome:
            t = tabelas.get(nome.lower())
        else:  # [@Col]: tabela que contem a propria celula
            t = next((t for t in tabelas.values() if t[0] == aqui
                      and t[1] - 1 <= pos[0] <= t[3] and t[2] <= pos[1] <= t[4]), None)
        if t is None:
            return None
        aba, r1, c1, r2, c2, cols = t
        pedidas = [x.strip().lstrip("@").lower() for x in re.split(r"[\[\],:]", dentro)]
        achadas = [c1 + cols.index(x) for x in pedidas if x in cols]
        if achadas:
            c1, c2 = min(achadas), max(achadas)
        return list(celulas(aba, r1, c1, r2, c2))

    def resolve(v, aqui, pos, abertos=()):
        """Operando de formula -> lista de celulas, ou None se nao der para seguir."""
        if re.match(r"^'?\[\d+\]", v):
            avisos[(aqui, "external file reference (not followed)")] += 1
            return []
        m = RE_TABELA.match(v)
        if m and "!" not in v:
            return tabela(m.group(1), m.group(2), aqui, pos)
        aba, alvo_txt = parte_aba(v)
        f = faixa(alvo_txt)
        if f is not None:
            aba = aqui if aba is None else abas.get(aba.lower())
            return None if aba is None else list(celulas(aba, *f))
        achado = nomes.get((aqui, v.lower())) or nomes.get((None, v.lower())) if aba is None else None
        if achado:
            nome, destino = achado
            usados.add(nome.lower())
            if "(" in destino or v.lower() in abertos:  # formula ou nome circular
                return None
            if re.match(r'^=?\s*(-?[\d.]+(E[+-]?\d+)?%?|".*"|TRUE|FALSE)\s*$', destino, re.I):
                return []  # constante (Taxa = 0.05): nada para seguir
            saida = []
            for parte in destino.split(","):
                achou = resolve(parte.strip(), aqui, pos, abertos + (v.lower(),))
                if achou is None:
                    return None
                saida += achou
            return saida
        if aba is None and v.lower() in tabelas:
            return tabela(v, "", aqui, pos)
        return None

    memo = {}
    fila, vistos = list(celulas(*alvo[:1] + (alvo[1], alvo[2], alvo[1], alvo[2]))) or [alvo], set()
    while fila:
        k = fila.pop()
        if k in vistos:
            continue
        vistos.add(k)
        f = dados[k[0]][0].get(k[1:])
        if not f or f[0] == "d":
            continue
        try:
            itens = Tokenizer(f[1]).items
        except Exception:  # ex.: B1# (spill) o Tokenizer nao entende
            chave = (k[0], "formula not parsed, refs not followed")
            avisos[chave] += 1
            exemplo.setdefault(chave, ref(*k[1:]))
            continue
        for t in itens:
            if t.type == Token.OPERAND and t.subtype == Token.RANGE:
                chave_memo = (t.value, k[0])
                if "[" in t.value:  # [@Col] depende da linha: sem cache
                    achou = resolve(t.value, k[0], k[1:])
                elif chave_memo in memo:
                    achou = memo[chave_memo]
                else:
                    achou = memo[chave_memo] = resolve(t.value, k[0], k[1:])
                if achou is None:
                    chave = (k[0], "unresolved reference %s" % t.value)
                    avisos[chave] += 1
                    exemplo.setdefault(chave, ref(*k[1:]))
                else:
                    fila.extend(x for x in achou if x not in vistos)
    cadeia = defaultdict(set)
    for aba, r, c in vistos:
        cadeia[aba].add((r, c))
    linhas = ["- '%s': %s%s%s" % (aba, txt, " in %s" % exemplo[(aba, txt)] if (aba, txt) in exemplo else "",
                                   " (x%d)" % n if n > 1 else "") for (aba, txt), n in sorted(avisos.items())]
    return cadeia, linhas, usados


def ler_abas(wb, wb_v):
    """[(aba, formulas, valores, cache)], tamanho do dump celula a celula, True se nenhuma formula tem valor salvo."""
    abas, bruto_chars, sem_cache = [], 0, True
    for ws in wb.worksheets:
        cache = {}
        for linha in wb_v[ws.title].iter_rows():
            for c in linha:
                if c.value is not None:
                    cache[(c.row, c.column)] = c.value
        formulas, valores = {}, {}
        for linha in ws.iter_rows():
            for c in linha:
                if c.value is None:
                    continue
                f = formula_de(c.value)
                pos = (c.row, c.column)
                if f is None:
                    valores[pos] = c.value
                    bruto_chars += len("%s\t%s" % (ref(*pos), fmt(c.value))) + 1
                else:
                    if f[0] == "f":
                        f = f + (relativa(f[1], c.row, c.column),)
                    formulas[pos] = f
                    bruto_chars += len("%s\t%s\t%s" % (ref(*pos), f[1] if len(f) > 1 else "TABLE()",
                                                      cache.get(pos))) + 1
        if any(pos in cache for pos in formulas):
            sem_cache = False
        abas.append((ws, formulas, valores, cache))

    return abas, bruto_chars, sem_cache


def converter(fonte, alvo=None):
    """Retorna (texto, numeros). fonte = caminho ou arquivo aberto.
    alvo = "Aba!B5": so a cadeia de precedentes dessa celula."""
    if hasattr(fonte, "seek"):
        fonte.seek(0)
    wb = load_workbook(fonte)
    if hasattr(fonte, "seek"):
        fonte.seek(0)
    wb_v = load_workbook(fonte, read_only=True, data_only=True)

    avisos, mapa, corpo = [], [], []
    falhas, total_formulas = [], 0
    ja_vistas = []  # (nome, esperado) das abas anteriores
    s_form = s_dados = 0

    abas, bruto_chars, sem_cache = ler_abas(wb, wb_v)

    cadeia, contexto, todas = None, None, 0
    if alvo is not None:
        alvo_pos = acha_alvo(wb, alvo)
        todas = sum(1 for _, fs, _, _ in abas for f in fs.values() if f[0] in "fa")
        cadeia, avisos_cadeia, usados = precedentes(wb, {ws.title: (f, v) for ws, f, v, _ in abas}, alvo_pos)
    for ws, formulas, valores, cache in abas:
        if cadeia is not None:
            meus = cadeia.get(ws.title)
            if not meus:
                continue
            contexto = (formulas, valores)  # rotulos olham a aba inteira
            formulas = {p: f for p, f in formulas.items() if p in meus}
            valores = {p: v for p, v in valores.items() if p in meus}

        textos = [f[1] for f in formulas.values() if f[0] in "fa"]
        total_formulas += len(textos)
        meu = esperado(formulas)

        # avisos da aba
        n_ref = sum(1 for t in textos if "#REF!" in t)
        n_erro = sum(1 for p in formulas if isinstance(cache.get(p), str) and cache[p] in ERROS)
        n_dinamica = sum(1 for t in textos if re.search(r"\b(OFFSET|INDIRECT)\(", t))
        n_tabela_dados = sum(1 for f in formulas.values() if f[0] == "d")
        nome = "'%s'" % ws.title
        if n_ref:
            avisos.append("- %s: %d formulas contain #REF! (broken reference)" % (nome, n_ref))
        if n_erro:
            avisos.append("- %s: %d formulas show an error value (#N/A, #DIV/0!...)" % (nome, n_erro))
        if n_dinamica:
            avisos.append("- %s: %d formulas use OFFSET/INDIRECT (reference decided at run time)" % (nome, n_dinamica))
        if n_tabela_dados:
            avisos.append("- %s: %d data-table (what-if) cells, not shown as formulas" % (nome, n_tabela_dados))

        # mapa
        usa, externo = Counter(), 0
        for t in textos:
            achadas = set()
            for q, s in RE_ABA.findall(t):
                achadas.add(q.replace("''", "'") if q else s)
            for a in achadas:
                if a != ws.title:
                    usa[a] += 1
            externo += bool(re.search(r"\[\d+\]", t))
        partes = ["- %s %s" % (nome, ws.dimensions)]
        if ws.sheet_state != "visible":
            partes.append("HIDDEN")
        partes.append("%d formulas" % len(textos))
        if usa:
            partes.append("uses " + ", ".join("'%s' (%d)" % kv for kv in usa.most_common()))
        if externo:
            partes.append("%d formulas point to other files" % externo)
        mapa.append(", ".join(partes))
        for t in ws.tables.values():
            colunas = [col.name for col in t.tableColumns]
            mapa.append("  table %s %s columns: %s" % (t.displayName, t.ref, ", ".join(colunas)))

        # formulas
        linhas_f = secao_formulas(formulas, valores, "blank" if cadeia is None else "not in chain", contexto, cache)
        if not confere(linhas_f, meu):
            falhas.append(ws.title)
            linhas_f = por_celula(formulas)
        for base_nome, base in ja_vistas:
            iguais = sum(1 for k, v in meu.items() if base.get(k) == v)
            uniao = len(set(meu) | set(base))
            if uniao and iguais / float(uniao) >= RATIO_DUPLICADA:
                dif = diferenca(formulas, meu, base_nome, base)
                if len("\n".join(dif)) < len("\n".join(linhas_f)) and confere(dif, meu, base):
                    linhas_f = dif
                    break
        ja_vistas.append((ws.title, meu))

        linhas_d = secao_dados(valores)
        bloco = ["", "# Sheet %s" % nome]
        if linhas_f:
            bloco += ["## Formulas"] + linhas_f
        if linhas_d:
            bloco += ["## Data"] + linhas_d
        s_form += tokens("\n".join(linhas_f))
        s_dados += tokens("\n".join(linhas_d))
        corpo += bloco

    wb_v.close()  # read_only segura o arquivo aberto
    if tem_vba(fonte):
        avisos.insert(0, "- File has VBA macros: some values may be written by a macro, not by formulas.")
    if total_formulas and sem_cache:
        avisos.insert(0, "- No calculated values saved in the file (never opened in Excel?): error checks are blind.")
    if falhas:
        avisos.insert(0, "- WARNING: compact view failed the self-check on %d sheet(s): %s. "
                         "Those sheets list one formula per cell." % (len(falhas), ", ".join(falhas)))
    nomes = nomes_definidos(wb)
    if cadeia is not None:
        avisos += avisos_cadeia
        nomes = [n for n in nomes if n.split(" = ")[0].lower() in usados]
    if nomes:
        mapa += ["Defined names:"] + ["  " + n for n in nomes]

    cabeca = [
        "# excel_ai dump",
        "How to read formulas: '<range> <formula>' gives the formula of the first cell;",
        "the rest of the range is the same formula copied (Excel fill), refs shifting as usual.",
        "'[label]' after the range is the nearest text header (above a column, left of a row): a hint, not part of the formula.",
        "'(v X)' is the value Excel last calculated for the first cell of the range: a hint, not part of the formula.",
        "'except <cell> ...' overrides one cell inside the range above it.",
        "'same formulas as X' means the sheet is a copy of X except the listed cells.",
    ]
    if cadeia is not None:
        aba, r, c = alvo_pos
        valor = next(cc for w, _, _, cc in abas if w.title == aba).get((r, c))
        cabeca += [
            "Precedents of '%s'!%s (cached value: %s): only the %d cells it depends on, in %d sheet(s)%s."
            % (aba, ref(r, c), "none" if valor is None else fmt(valor), sum(map(len, cadeia.values())), len(cadeia),
               "; some refs could not be followed, see Warnings" if avisos_cadeia else ""),
            "'except <cell> not in chain' = that cell is not used by the target.",
        ]
    texto_avisos = ["", "## Warnings"] + (avisos or ["- none"])
    texto_mapa = ["", "## Map"] + mapa
    texto = "\n".join(cabeca + texto_avisos + texto_mapa + corpo) + "\n"
    numeros = {
        "bruto": max(bruto_chars - 1, 0) // 4,  # = tokens() do dump juntado com "\n"
        "novo": tokens(texto),
        "avisos": tokens("\n".join(texto_avisos)),
        "mapa": tokens("\n".join(texto_mapa)),
        "formulas": s_form,
        "dados": s_dados,
        "total_formulas": total_formulas,
        "todas_formulas": todas,
        "falhas": len(falhas),
    }
    return texto, numeros


LIMITE_CONSTANTES = 200  # linhas de constantes por aba no --diff


def _igual(a, b):
    """Mesmo valor (ruido de ponto flutuante abaixo de 1e-9 nao conta)."""
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))
    return fmt(a) == fmt(b)


def _abrir(fonte):
    """(wb, abas) de uma planilha, para o --diff."""
    if hasattr(fonte, "seek"):
        fonte.seek(0)
    wb = load_workbook(fonte)
    if hasattr(fonte, "seek"):
        fonte.seek(0)
    wb_v = load_workbook(fonte, read_only=True, data_only=True)
    abas = ler_abas(wb, wb_v)[0]
    wb_v.close()
    return wb, {ws.title: (formulas, valores, cache) for ws, formulas, valores, cache in abas}


def _secao_diff(titulo, formulas, contexto, cache, falhas, aba):
    """Bloco compacto de um conjunto de formulas, com a conferencia de sempre."""
    if not formulas:
        return []
    linhas = secao_formulas(formulas, {}, "unchanged", contexto, cache)
    if not confere(linhas, esperado(formulas)):
        falhas.append("%s (%s)" % (aba, titulo))
        linhas = por_celula(formulas)
    return ["## " + titulo] + linhas


def converter_diff(antiga, nova):
    """Retorna (texto, numeros): so o que mudou da planilha antiga para a nova."""
    wb_a, abas_a = _abrir(antiga)
    wb_n, abas_n = _abrir(nova)
    falhas, corpo, resumo = [], [], []
    n_mudou = n_antes = n_const = n_iguais = 0
    vazio = ({}, {}, {})

    for nome in list(abas_n) + [t for t in abas_a if t not in abas_n]:
        if nome not in abas_a or nome not in abas_n:
            fs = (abas_n if nome in abas_n else abas_a)[nome][0]
            qtd = sum(1 for f in fs.values() if f[0] in "fa")
            resumo.append("- sheet '%s' %s (%d formulas)" % (nome, "added" if nome in abas_n else "removed", qtd))
            continue
        fa, va, ca = abas_a[nome]
        fn, vn, cn = abas_n[nome]
        ea, en = esperado(fa), esperado(fn)
        mudou = {pos: f for pos, f in fn.items() if f[0] in "fa" and ea.get(ref(*pos)) != en[ref(*pos)]}
        antes = {pos: f for pos, f in fa.items() if f[0] in "fa" and en.get(ref(*pos)) != ea[ref(*pos)]}
        const = []
        va = {q: x for q, x in va.items() if not (isinstance(x, str) and not x.strip())}  # "" = vazio
        vn = {q: x for q, x in vn.items() if not (isinstance(x, str) and not x.strip())}
        for pos in sorted(set(va) | set(vn)):
            a, n = va.get(pos), vn.get(pos)
            if pos in va and pos in vn:
                if not _igual(a, n):
                    const.append("%s %s -> %s" % (ref(*pos), fmt(a), fmt(n)))
            elif pos in vn:
                const.append("%s added %s" % (ref(*pos), fmt(n)))
            else:
                const.append("%s removed (was %s)" % (ref(*pos), fmt(a)))
        if not (mudou or antes or const):
            n_iguais += 1
            continue
        n_mudou += len(mudou)
        n_antes += len(antes)
        n_const += len(const)
        bloco = ["", "# Sheet '%s'" % nome]
        bloco += _secao_diff("Formulas new or changed", mudou, (fn, vn), cn, falhas, nome)
        bloco += _secao_diff("Formulas before (changed or removed cells)", antes, (fa, va), ca, falhas, nome)
        if const:
            extra = len(const) - LIMITE_CONSTANTES
            bloco += ["## Constants"] + const[:LIMITE_CONSTANTES] + (["... %d more" % extra] if extra > 0 else [])
        corpo += bloco

    nomes_a, nomes_n = set(nomes_definidos(wb_a)), set(nomes_definidos(wb_n))  # "NOME = destino"; destino novo = removido + novo
    nomes = ["removed: " + n for n in sorted(nomes_a - nomes_n)] + ["added: " + n for n in sorted(nomes_n - nomes_a)]

    cabeca = [
        "# excel_ai diff (old -> new)",
        "Only what changed. '## Formulas new or changed' is the NEW version, same format as the normal dump",
        "('<range> [label] (v value) <formula>', 'except' = cell inside the range that did not change).",
        "'## Formulas before' is the OLD formula of the cells that changed or disappeared.",
        "'A1 old -> new' is a constant that changed.",
        "Summary: %d formulas new or changed, %d old formulas replaced or removed, %d constants changed, %d sheet(s) without changes."
        % (n_mudou, n_antes, n_const, n_iguais),
    ]
    if falhas:
        cabeca.append("WARNING: compact view failed the self-check in: %s. Those parts list one formula per cell." % ", ".join(falhas))
    if resumo:
        cabeca += ["", "## Sheets"] + resumo
    if nomes:
        cabeca += ["", "## Defined names"] + nomes
    if not (corpo or resumo or nomes):
        cabeca.append("No changes found.")
    texto = "\n".join(cabeca + corpo) + "\n"
    return texto, {"novo": tokens(texto), "formulas": n_mudou, "antes": n_antes, "constantes": n_const,
                   "falhas": len(falhas)}


def main_diff(antiga, nova):
    for caminho in (antiga, nova):
        if not os.path.isfile(caminho):
            sys.stderr.write("erro: arquivo nao encontrado: %s\n" % ascii(caminho))
            return 2
    try:
        texto, n = converter_diff(antiga, nova)
    except Exception as e:
        sys.stderr.write("erro: nao consegui ler a planilha (%s: %s)\n" % (type(e).__name__, ascii(str(e))))
        return 2
    limpo = lambda c: re.sub(r"[^A-Za-z0-9]", "_", os.path.splitext(os.path.basename(c))[0])
    destino = os.path.join(os.path.dirname(os.path.abspath(nova)), "%s.vs_%s.diff.ai.txt" % (limpo(nova), limpo(antiga)))
    with io.open(destino, "w", encoding="utf-8", newline="\n") as arq:
        arq.write(texto)
    print("Tokens do diff:  %10s" % "{:,}".format(n["novo"]))
    print("Formulas novas/mudadas: %s | antigas trocadas/removidas: %s | constantes: %s" % (
        n["formulas"], n["antes"], n["constantes"]))
    if n["falhas"]:
        print("Conferencia das formulas: FALHOU em %d parte(s); elas sairam celula a celula." % n["falhas"])
    else:
        print("Conferencia das formulas: 100%%")
    print("Arquivo gerado: %s" % os.path.basename(destino).encode("ascii", "backslashreplace").decode())
    return 1 if n["falhas"] else 0


def main(caminho, alvo=None):
    if not os.path.isfile(caminho):
        sys.stderr.write("erro: arquivo nao encontrado: %s\n" % ascii(caminho))
        return 2
    try:
        texto, n = converter(caminho, alvo)
    except ValueError as e:
        sys.stderr.write("erro: %s\n" % ascii(str(e)))
        return 2
    except Exception as e:
        sys.stderr.write("erro: nao consegui ler a planilha (%s: %s)\n" % (type(e).__name__, ascii(str(e))))
        return 2
    sufixo = ""
    if alvo is not None:
        aba, cel = parte_aba(alvo.strip())
        sufixo = "." + re.sub(r"[^A-Za-z0-9]", "_", "%s_%s" % (aba or "ativa", cel.replace("$", "").upper()))
    destino = os.path.splitext(caminho)[0] + sufixo + ".ai.txt"
    with io.open(destino, "w", encoding="utf-8", newline="\n") as arq:
        arq.write(texto)
    print("Tokens hoje (celula a celula): %10s" % "{:,}".format(n["bruto"]))
    print("Tokens com excel_ai:           %10s" % "{:,}".format(n["novo"]))
    print("  avisos %s | mapa %s | formulas %s | dados %s" % tuple(
        "{:,}".format(n[k]) for k in ("avisos", "mapa", "formulas", "dados")))
    if n["novo"]:
        print("Reducao: %.1fx" % (n["bruto"] / float(n["novo"])))
    if alvo is not None:
        print("Cadeia: %s de %s formulas da planilha (%.1f%%)" % (
            "{:,}".format(n["total_formulas"]), "{:,}".format(n["todas_formulas"]),
            100.0 * n["total_formulas"] / max(n["todas_formulas"], 1)))
    if n["falhas"]:
        print("Conferencia das formulas: FALHOU em %d aba(s); essas abas sairam celula a celula." % n["falhas"])
    else:
        print("Conferencia das formulas: 100%% (%s formulas sem perda)" % "{:,}".format(n["total_formulas"]))
    print("Arquivo gerado: %s" % os.path.basename(destino).encode("ascii", "backslashreplace").decode())
    return 1 if n["falhas"] else 0


def main_pasta(pasta):
    """Converte cada planilha da pasta com main(); um arquivo ruim nao para os outros."""
    nomes = sorted(n for n in os.listdir(pasta)
                   if n.lower().endswith((".xlsx", ".xlsm")) and not n.startswith("~$")
                   and os.path.isfile(os.path.join(pasta, n)))
    if not nomes:
        sys.stderr.write("erro: nenhuma planilha .xlsx/.xlsm na pasta: %s\n" % ascii(pasta))
        return 2
    codigos = []
    for n in nomes:
        print("== %s ==" % n.encode("ascii", "backslashreplace").decode())
        codigos.append(main(os.path.join(pasta, n)))
    print("Pasta: %d planilha(s), %d com problema" % (len(nomes), sum(1 for c in codigos if c)))
    return max(codigos)


def _fixture():
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    p = wb.active
    p.title = "Premissas"
    p["A1"], p["B1"] = "Taxa", 0.012
    p["A2"], p["B2"] = "Principal", 1000000
    p["A3"], p["B3"] = "Carencia", 6

    f = wb.create_sheet("Fluxo")
    f.append(["Parcela", "Saldo", "Taxa", "Juros"])
    for r in range(2, 102):
        f["A%d" % r] = r - 1
        f["B%d" % r] = "=Premissas!$B$2-A%d" % r
        f["C%d" % r] = "=_xlfn.XLOOKUP(\"Taxa\",Premissas!$A:$A,Premissas!$B:$B)"
        f["D%d" % r] = "=B%d*C%d" % (r, r)
    f["D50"] = "=B50*0"   # formula diferente no meio
    f["D60"] = 7          # numero digitado no meio
    f["D70"] = None       # vazio no meio
    f["D80"], f["D81"] = "=1", "=2"  # duas diferentes juntas
    f["F2"] = ArrayFormula("F2:F4", "=B2:B4*2")
    f["G2"] = "=OFFSET(A1,1,0)"
    f["H2"] = "=#REF!*2"
    f["I2"] = "=SUM(Tabela1[Valor])"
    f["J2"] = "='Plan Extra 2'!A1+TAXA2024"

    m = wb.create_sheet("Meses")
    for c in range(2, 14):
        m.cell(row=1, column=c, value=c * 10)
        m.cell(row=2, column=c, value="=%s1*2" % get_column_letter(c))

    t = wb.create_sheet("Plan Extra 2")
    t.append(["Nome", "Valor"])
    for i in range(3):
        t.append(["item%d" % i, i])
    t.add_table(Table(displayName="Tabela1", ref="A1:B4"))
    t.sheet_state = "hidden"

    copia = wb.copy_worksheet(f)
    copia.title = "Fluxo B"
    copia["D30"] = "=B30*C30*2"
    copia["A2"] = 999

    nome = DefinedName("TAXA2024", attr_text="Premissas!$B$1")
    try:
        wb.defined_names.add(nome)
    except AttributeError:
        wb.defined_names.append(nome)
    buf = io.BytesIO()
    wb.save(buf)
    return buf


SKILL_MD = r'''---
name: excel-ai
description: Read Excel workbooks (.xlsx/.xlsm) efficiently by converting them to compact, lossless text before analysing them. Use whenever the user asks about a spreadsheet, its formulas, calculation logic, cross-sheet references (XLOOKUP/PROCX), financial calculators (IRR/TIR, duration, price schedules, provisioning/PDD), why a cell has a value, what changed between two versions of a workbook, or wants a whole folder of workbooks read. Triggers - planilha, Excel, xlsx, xlsm, formula, PROCX, TIR, calculadora, diff de versoes.
license: MIT
---

# excel-ai: read spreadsheets as compact text

`scripts/excel_ai.py` turns a workbook into a short text file that keeps every formula and constant that matters and drops formatting noise. A 1.6M-token cell-by-cell dump of a large financial calculator becomes about 33k tokens (about 48x smaller). Formulas are verified lossless: the text is expanded back and compared with the original.

Never read an `.xlsx`/`.xlsm` by unzipping it or dumping cells with openpyxl yourself. Run the tool and read its output.

## Setup (once)

```
pip install openpyxl
```

Python 3.8+. One file, no network, nothing leaves the machine. The script is at `scripts/excel_ai.py` in this skill folder. Run it with that path (written below as `excel_ai.py`).

## Pick the mode

| The user wants... | Run |
|---|---|
| To understand or audit one workbook | `python excel_ai.py book.xlsx` |
| One number explained (final IRR, a total, a provision) | `python excel_ai.py book.xlsx --celula "Sheet!B5"` |
| What changed between two versions | `python excel_ai.py old.xlsx new.xlsx --diff` |
| Every workbook in a folder | `python excel_ai.py folder` |

Prefer `--celula` over the full dump whenever the question is about one result: it keeps only the cells that result depends on (for a real calculator, about 6% of the formulas and 2k tokens instead of 33k). Prefer `--diff` over reading two dumps.

Output files are written next to the input:

- full: `book.ai.txt`
- precedents: `book.<Sheet>_<CELL>.ai.txt` (non-alphanumerics become `_`)
- diff: `new.vs_old.diff.ai.txt`
- folder: one `.ai.txt` per workbook (not recursive; `~$*` lock files skipped)

After running, read the generated `.ai.txt` file (not the workbook). The console also prints tokens before/after and the self-check result.

## `--celula` details

- Reference forms: `'Sheet name'!B5`, `Sheet!B5`, `Sheet!$b$5`, or `B5` (active sheet). One cell only, never a range.
- It follows the chain through other sheets, ranges, whole columns/rows (limited to existing cells), defined names, table references (`Table[Col]`, `[@Col]`) and array formulas.
- Not followed, but listed in Warnings: `OFFSET`/`INDIRECT` results, external-file references, data-table (what-if) inputs, names that are not plain ranges, `#REF!`, and LET/LAMBDA variables. If Warnings says "some refs could not be followed", tell the user the chain may be incomplete and offer the full dump.
- Lines `except <cell> not in chain` are gaps inside a block: formulas that exist but are not used by the target.
- The header shows the target's cached value (what Excel last calculated).

## `--diff` details

- Sheets match by exact name. Added/removed sheets are listed with formula counts.
- `## Formulas new or changed` is the NEW version in the normal compact format; `## Formulas before` is the OLD formula of changed or removed cells. Constants appear as `A1 old -> new`, `A1 added X`, `A1 removed (was X)` (200 lines per sheet max). Defined names appear as `added:`/`removed:` (a changed destination is removed + added).
- A pure fill-down with the same relative formula is not a change. Float noise below 1e-9 and empty strings are ignored.
- Inserting rows or columns shows up as many changes (no move detection). Data-table cells are not compared. Say so when the diff looks huge.
- "No changes found." means formulas, constants and names are identical.

## Reading the output

Read `references/format.md` for the full legend. Short version:

- `## Warnings` first: `#REF!`, error values, OFFSET/INDIRECT, macros, uncalculated files. Mention relevant ones to the user.
- `## Map`: each sheet (dimensions, HIDDEN, formula count, which sheets it uses), tables, defined names.
- Per sheet `## Formulas`: `E2:E361 [Juros] (v 12.5) =C2*G2` means the formula of the first cell, copied down/across the range with references shifting. `[Juros]` is the nearest text header (a hint). `(v 12.5)` is the value Excel last calculated for the first cell (a hint, blocks of 2+ cells only). `  except E9 =...` overrides one cell. `same formulas as 'X', except:` means a copy of sheet X with listed differences.
- Per sheet `## Data`: constants, with long number runs summarized as `[100 numbers] min .. max .. first: ..`.

Reason about the logic from the formulas: say what each block computes, which sheets feed it, and where inputs live. Use cached values `(v ...)` only as sanity checks; they may be stale.

## Exit codes and problems

- `0` ok. `1` the lossless self-check failed on some sheet (that sheet is listed one formula per cell, still correct, just larger). `2` bad arguments, missing file, unreadable workbook, or a folder with no workbooks.
- "No calculated values saved": the file was never opened in Excel, so cached values and error checks are missing.
- "File has VBA macros": some values may come from a macro, which this tool does not read.
- `.xls` (old format) and `.xlsb` are not supported; ask the user to save as `.xlsx`.
- Large workbooks take 15-40 seconds. Do not interrupt.

## Rules

1. Run the tool before answering any question about a workbook's contents.
2. Never paste real workbook data into public places (issues, PRs, public repos). Workbooks are usually confidential.
3. Do not edit the workbook with this tool; it is read-only.
4. If the user's question is about one result, start with `--celula`. If it is about differences between versions, use `--diff`.
5. When the output says a part could not be followed or verified, say it plainly instead of guessing.
'''

FORMAT_MD = r'''# excel_ai output format

## File layout

```
# excel_ai dump
<how-to-read lines>

## Warnings
- ...

## Map
- 'Sheet' A1:J101, 303 formulas, uses 'Other' (200)
  table Tabela1 A1:B4 columns: Nome, Valor
Defined names:
  TAXA2024 = Premissas!$B$1

# Sheet 'Fluxo'
## Formulas
...
## Data
...
```

## Formula lines

| Line | Meaning |
|---|---|
| `D2:D101 =B2*C2` | Formula of D2; D3..D101 are the same formula filled down (relative references shift as in Excel). |
| `C2:E2 =C1*2` | Same, filled across. |
| `D2:D101 [Juros] =B2*C2` | `[Juros]` is the nearest text header: above for columns, left for rows, up to 2 empty cells away. A hint, not part of the formula. |
| `D2:D101 [Juros] (v 12.5) =B2*C2` | `(v 12.5)` is the cached value (last Excel result) of the FIRST cell. Only blocks of 2+ cells. Hint only. |
| `  except D30 =B30*C30*2` | D30 inside the block above has a different formula. |
| `  except D30 value 5` | D30 is a constant inside the block. |
| `  except D30 blank` | D30 is empty. In `--celula` mode: `not in chain` (exists but not used by the target). In a copy-sheet: `unchanged`. |
| `E1 array E1:E2 {=Z1:Z2*2}` | Array formula anchored at E1 covering E1:E2. Other cells of the range are not listed. |
| `same formulas as 'Fluxo', except:` | The sheet is a copy (80%+ equal) of that sheet; only differing cells follow. `A1 no formula` means the cell has no formula here. |

What-if data-table cells are not listed as lines; they only appear counted in Warnings.

Internal prefixes (`_xlfn.`) are removed: `_xlfn.XLOOKUP` is shown as `XLOOKUP` (PROCX in Portuguese Excel).

## Data lines

| Line | Meaning |
|---|---|
| `A1 Parcela \| B1 Saldo` | Constants on a row, `cell value` separated by `\|`. |
| `A2:A101 [100 numbers] min 1 max 100 first: 1, 2, 3` | Long run of numbers/dates summarized. Runs of up to 10 stay complete. |
| `A2:A69 [68 texts, 68 distinct] first: a, b, c` | Long run of texts (over 50) summarized. Shorter runs stay complete. |
| `A5:A9 x (x5)` | Same text repeated. |

## Warnings you may see

- `N formulas contain #REF!` broken reference.
- `N formulas show an error value (#N/A, #DIV/0!...)` from cached values.
- `N formulas use OFFSET/INDIRECT` target decided at run time; not followable statically.
- `N data-table (what-if) cells`.
- `File has VBA macros`.
- `No calculated values saved in the file`.
- `compact view failed the self-check on N sheet(s)` those sheets are one formula per cell.
- `--celula` only: `unresolved` / `formula not parsed` refs that could not be followed.

## Diff file

```
# excel_ai diff (old -> new)
Summary: N formulas new or changed, M old formulas replaced or removed, K constants changed, S sheet(s) without changes.
## Sheets          (added/removed sheets)
## Defined names   (added:/removed:)
# Sheet 'X'
## Formulas new or changed
## Formulas before (changed or removed cells)
## Constants
```
'''


def instalar_skill(destino=None, forcar=False):
    """Grava a skill do Kiro (SKILL.md, references/format.md e uma copia deste script). Pula se ja existe."""
    destino = destino or os.path.expanduser("~/.kiro/skills/excel-ai")
    if os.path.exists(destino) and not forcar:
        print("Skill ja instalada em %s (use --forcar para sobrescrever)" % destino)
        return 0
    for rel, texto in (("SKILL.md", SKILL_MD), ("references/format.md", FORMAT_MD)):
        caminho = os.path.join(destino, rel)
        os.makedirs(os.path.dirname(caminho), exist_ok=True)
        with open(caminho, "w", encoding="utf-8") as f:
            f.write(texto)
    os.makedirs(os.path.join(destino, "scripts"), exist_ok=True)
    shutil.copyfile(os.path.abspath(__file__), os.path.join(destino, "scripts", "excel_ai.py"))
    print("Skill instalada em %s" % destino)
    return 0


def autoteste():
    # chave relativa: mesmo padrao na coluna inteira
    assert relativa("=B2*C2", 2, 5) == "=R[0]C[-3]*R[0]C[-2]"
    assert relativa("=B3*C3", 3, 5) == relativa("=B2*C2", 2, 5)
    assert relativa("=B2*$A$1", 2, 3) == "=R[0]C[-1]*R1C1"
    assert relativa("=$B2*B$1", 2, 3) == "=R[0]C2*R1C[-1]"
    assert relativa("='Plan 1'!B2", 2, 2) == "='Plan 1'!R[0]C[0]"
    assert relativa('=IF(A2="B2",1,0)', 2, 2) == '=IF(R[0]C[-1]="B2",1,0)'
    assert relativa("=Tabela1[CDI2024]", 2, 2) == "=Tabela1[CDI2024]"
    assert relativa("=Tabela1[[#This Row],[A1]]", 2, 2) == "=Tabela1[[#This Row],[A1]]"
    assert relativa("=TAXA2024*A2", 2, 2) == "=TAXA2024*R[0]C[-1]"
    assert relativa("=SUM(B:B)", 1, 3) == relativa("=SUM(C:C)", 1, 4)
    assert relativa("=SUM(2:2)", 2, 1) == relativa("=SUM(3:3)", 3, 1)
    assert relativa("=SUM($B:$B)", 1, 3) != relativa("=SUM($C:$C)", 1, 4)
    assert limpa("=_xlfn.XLOOKUP(A1,B:B,C:C)") == "=XLOOKUP(A1,B:B,C:C)"

    texto, n = converter(_fixture())
    secoes = texto.split("\n# Sheet ")
    fluxo = [s for s in secoes if s.startswith("'Fluxo'")][0]
    # bloco com excecoes: formula diferente, numero e vazio no meio
    assert "D2:D79 [Juros] =B2*C2" in fluxo, fluxo
    assert "  except D50 =B50*0" in fluxo
    assert "  except D60 value 7" in fluxo
    assert "  except D70 blank" in fluxo
    assert "D80 =1" in fluxo and "D81 =2" in fluxo and "D82:D101 =B82*C82" in fluxo  # formula acima: sem rotulo
    assert "C2:C101 [Taxa] =XLOOKUP(" in fluxo and "_xlfn" not in fluxo
    assert "F2 array F2:F4 {=B2:B4*2}" in fluxo
    assert "I2 =SUM(Tabela1[Valor])" in fluxo
    # meses nas colunas: agrupa na horizontal
    assert "B2:M2 =B1*2" in texto
    # copia de aba: so a diferenca, mas os dados ficam inteiros
    copia = [s for s in secoes if s.startswith("'Fluxo B'")][0]
    assert "same formulas as 'Fluxo', except:" in copia and "D30 =B30*C30*2" in copia
    # diferencas tambem agrupadas; celula igual a base no meio vira 'unchanged'
    dif = diferenca({(r, 1): ("f", "=X%d" % r, relativa("=X%d" % r, r, 1)) for r in range(2, 9)},
                    {"A%d" % r: "=X%d" % r for r in range(2, 9)}, "B",
                    {"A5": "=X5", "A9": "=Y"})
    assert dif == ["same formulas as 'B', except:", "A2:A8 =X2", "  except A5 unchanged", "A9 no formula"], dif
    assert expandir(dif, {"A5": "=X5", "A9": "=Y"}) == {"A%d" % r: "=X%d" % r for r in range(2, 9)}
    # textos iguais seguidos viram uma linha; coluna longa de textos vira resumo
    linhas = secao_dados({(r, 1): "INSS" for r in range(2, 8)})
    assert linhas == ["A2:A7 INSS (x6)"], linhas
    linhas = secao_dados({(r, 1): "cod%d" % r for r in range(2, 70)})
    assert linhas == ["A2:A69 [68 texts, 68 distinct] first: cod2, cod3, cod4"], linhas
    assert "A2:A101 [100 numbers] min 2 max 999 first: 999, 2, 3" in copia
    # dados: texto inteiro, sequencia longa de numeros resumida, premissas inteiras
    assert "A1 Parcela | B1 Saldo | C1 Taxa | D1 Juros" in fluxo
    assert "A2:A101 [100 numbers] min 1 max 100 first: 1, 2, 3" in fluxo
    assert "A1 Taxa | B1 0.012" in texto
    # mapa e avisos
    assert "- 'Fluxo' A1:J101, 303 formulas, uses 'Premissas' (200), 'Plan Extra 2' (1)" in texto
    assert "'Plan Extra 2' A1:B4, HIDDEN" in texto
    assert "  table Tabela1 A1:B4 columns: Nome, Valor" in texto
    assert "  TAXA2024 = Premissas!$B$1" in texto
    assert "- 'Fluxo': 1 formulas contain #REF!" in texto
    assert "- 'Fluxo': 1 formulas use OFFSET/INDIRECT" in texto
    assert "No calculated values saved" in texto
    assert "VBA" not in texto
    assert n["falhas"] == 0 and n["total_formulas"] > 0

    # rotulo: cabecalho acima (coluna) ou a esquerda (linha); pula ate 2 vazias
    fs = {(r, 2): ("f", "=A%d" % r, relativa("=A%d" % r, r, 2)) for r in range(5, 8)}
    assert blocos(fs, {(2, 2): "Juros"}, True)[0] == "B5:B7 [Juros] =A5"
    assert blocos(fs, {(1, 2): "Juros"}, True)[0] == "B5:B7 =A5"  # 3 vazias: longe demais
    assert blocos(fs, {(4, 2): 12}, True)[0] == "B5:B7 =A5"  # numero nao e rotulo
    assert blocos(fs, {(4, 2): "x [a]  muito   comprido " * 3}, True)[0] == "B5:B7 [x a muito comprido x a muito c] =A5"
    fs[(4, 2)] = ("f", "=1", "=1")
    assert blocos(fs, {(3, 2): "Juros"}, True)[1] == "B5:B7 =A5"  # formula acima para a busca
    fh = {(3, c): ("f", "=%s1" % get_column_letter(c), relativa("=%s1" % get_column_letter(c), 3, c)) for c in range(3, 6)}
    assert blocos(fh, {(3, 1): "Total"}, False)[0] == "C3:E3 [Total] =C1"
    assert blocos(fh, {(3, 2): "%"}, False)[0] == "C3:E3 =C1"  # unidade nao e rotulo
    assert blocos(fh, {(3, 2): "x"}, False)[0] == "C3:E3 =C1"  # uma letra so nao e rotulo
    assert blocos(fh, {(3, 2): "#N/A"}, False)[0] == "C3:E3 =C1"  # erro nao e rotulo
    assert blocos(fh, {(3, 2): "  ", (3, 1): "Total"}, False)[0] == "C3:E3 [Total] =C1"  # espaco = vazio
    um = {(4, 3): ("f", "=A1", relativa("=A1", 4, 3))}
    assert blocos(um, {(4, 2): 5, (3, 3): "TIR"}, True) == ["C4 =A1"]  # numero a esquerda: sem rotulo
    assert blocos(um, {(3, 3): "TIR"}, True) == ["C4 [TIR] =A1"]  # nada a esquerda: olha acima
    assert blocos(um, {(4, 2): "Taxa", (3, 3): "TIR"}, True) == ["C4 [Taxa] =A1"]  # esquerda ganha
    assert expandir(["B5:B7 [Juros] =A5"]) == expandir(["B5:B7 =A5"])

    # valor calculado: so o 1o da faixa, depois do rotulo; sem cache ou None = sem dica
    fs = {(r, 2): ("f", "=A%d" % r, relativa("=A%d" % r, r, 2)) for r in range(5, 8)}
    assert blocos(fs, {(4, 2): "Juros"}, True, cache={(5, 2): 12.5, (6, 2): 99})[0] == "B5:B7 [Juros] (v 12.5) =A5"
    assert blocos(fs, {}, True, cache={(5, 2): 3})[0] == "B5:B7 (v 3) =A5"
    assert blocos({(5, 2): fs[(5, 2)]}, {}, True, cache={(5, 2): 3})[0] == "B5 =A5"  # celula sozinha: sem dica
    assert blocos(fs, {}, True, cache={(6, 2): 3})[0] == "B5:B7 =A5"  # 1a celula sem valor
    assert blocos(fs, {}, True, cache={(5, 2): None})[0] == "B5:B7 =A5"
    assert blocos(fs, {}, True, cache={(5, 2): "a)b\nc  " + "x" * 30})[0] == "B5:B7 (v ab c xxxxxxxxxxxxxxx) =A5"
    assert blocos(fs, {}, True, cache={(5, 2): "#N/A"})[0] == "B5:B7 (v #N/A) =A5"
    assert blocos(fs, {}, True, cache={(5, 2): True})[0] == "B5:B7 (v TRUE) =A5"
    assert blocos(fs, {}, True)[0] == "B5:B7 =A5"
    fx = dict(fs)
    fx[(6, 2)] = ("f", "=Z1", relativa("=Z1", 6, 2))
    assert not any("(v" in l for l in blocos(fx, {}, True, cache={(5, 2): 1, (6, 2): 2}) if l.startswith("  except"))
    assert expandir(["B5:B7 [Juros] (v 12.5) =A5"]) == expandir(["B5:B7 =A5"])
    assert expandir(["B5 (v 1) =A5"]) == {"B5": "=A5"}
    assert not confere(["B5:B7 (v 1) =A5"], {"B5": "=A5", "B6": "=A6", "B7": "=A8"})
    assert not confere(["B5:B7 (v 1 =A5"], {"B5": "=A5", "B6": "=A6", "B7": "=A7"})
    certo = {"D%d" % r: "=B%d*C%d" % (r, r) for r in range(2, 11)}
    assert confere(["D2:D10 [J] =B2*C2"], certo) and not confere(["D2:D10 [J] =B2*C3"], certo)
    assert not confere(["D2:D10 [J =B2*C2"], certo)  # rotulo mal fechado nao passa

    # conferencia pega uma linha estragada
    linhas = ["D2:D10 =B2*C2"]
    certo = {"D%d" % r: "=B%d*C%d" % (r, r) for r in range(2, 11)}
    assert expandir(linhas) == certo
    assert expandir(["D2:D10 =B2*C3"]) != certo
    assert not confere(["linha que nao e formula"], certo)

    # formula com quebra de linha (Alt+Enter) sai numa linha so e confere
    wb = Workbook()
    wb.active["A1"], wb.active["B1"] = 1, "=IF(A1>0,\n1,2)"
    buf = io.BytesIO()
    wb.save(buf)
    texto, n = converter(buf)
    assert "B1 =IF(A1>0, 1,2)" in texto and n["falhas"] == 0

    # aba copiada com matriz de faixa diferente nao passa como igual
    from openpyxl.worksheet.formula import ArrayFormula
    wb = Workbook()
    for nome, faixa in (("S1", "C1:C3"), ("S2", "C1:C2")):
        ws = wb.create_sheet(nome)
        for r in range(1, 30):
            ws["A%d" % r] = "=B%d" % r
        ws["C1"] = ArrayFormula(faixa, "=A1:A3*2")
    buf = io.BytesIO()
    wb.save(buf)
    texto, n = converter(buf)
    assert "C1 array C1:C2 {=A1:A3*2}" in texto.split("# Sheet 'S2'")[1] and n["falhas"] == 0

    # plano B: se a conferencia falha, a aba sai celula a celula, com aviso, e main() sai com 1
    import contextlib
    import tempfile
    pasta = tempfile.mkdtemp()
    caminho = os.path.join(pasta, "fix.xlsx")
    with open(caminho, "wb") as arq:
        arq.write(_fixture().getvalue())
    with contextlib.redirect_stdout(io.StringIO()):
        assert main(caminho) == 0 and os.path.isfile(os.path.join(pasta, "fix.ai.txt"))
        original = globals()["expandir"]
        globals()["expandir"] = lambda linhas, base=None: {}
        try:
            texto, n = converter(caminho)
            assert main(caminho) == 1
        finally:
            globals()["expandir"] = original
    # 3 abas com formulas (Fluxo, Meses, Fluxo B); abas sem formula nao tem o que conferir
    assert n["falhas"] == 3 and "WARNING: compact view failed the self-check on 3 sheet(s)" in texto
    assert "D3 =B3*C3" in texto  # uma linha por formula
    with contextlib.redirect_stderr(io.StringIO()):
        assert main(os.path.join(pasta, "nao_existe.xlsx")) == 2
    # Translator do openpyxl apagava ABA!#REF!; tradutor() protege
    assert tradutor("=XLOOKUP(A2,Base!#REF!,'P 1'!#REF!)+#REF!", "B2")("B3") == "=XLOOKUP(A3,Base!#REF!,'P 1'!#REF!)+#REF!"

    # cadeia de precedentes (--celula): so o que a celula usa, entre abas
    texto, n = converter(_fixture(), "Fluxo!D5")
    assert "Precedents of 'Fluxo'!D5" in texto, texto[:400]
    fl = texto.split("# Sheet 'Fluxo'")[1].split("# Sheet ")[0]
    assert "D5 =B5*C5" in fl and "B5 =Premissas!$B$2-A5" in fl and "C5 =XLOOKUP(" in fl, fl
    assert "D6" not in fl and "Meses" not in texto and "Fluxo B" not in texto
    assert "A3 Carencia | B3 6" in texto  # coluna inteira do XLOOKUP entra
    assert n["falhas"] == 0 and n["total_formulas"] == 3
    texto, n = converter(_fixture(), "fluxo!$j$2")  # nome definido + outra aba
    assert "TAXA2024 = Premissas!$B$1" in texto and "A1 Nome" in texto and "B1 0.012" in texto
    texto, n = converter(_fixture(), "Fluxo!I2")  # tabela: so a coluna Valor
    extra = texto.split("# Sheet 'Plan Extra 2'")[1]
    assert "B2" in extra and "item0" not in extra, extra
    for ruim in ("Nada!A1", "Fluxo!A1:B2", "Fluxo!ZZZZ1"):
        try:
            converter(_fixture(), ruim)
            assert False, ruim
        except ValueError:
            pass
    from openpyxl.worksheet.table import Table
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "Valor"
    for r in range(2, 13):
        ws["A%d" % r], ws["B%d" % r] = "=B%d*2" % r, r
    w2 = wb.create_sheet("T2")
    w2["D1"], w2["E1"] = "=E1", "=D1"  # ciclo
    w2["F1"], w2["G1"] = "=Nope!A1", '=INDIRECT("A1")'
    w2["H1"] = ArrayFormula("H1:H3", "=B2:B4*2")
    w2["I1"] = "=H2"  # meio da matriz puxa a ancora
    w2["K1"], w2["L1"], w2["K2"], w2["K3"] = "Qtd", "Dobro", 7, 8
    w2["L2"], w2["L3"] = "=[@Qtd]*2", "=[@Qtd]*2"
    w2.add_table(Table(displayName="T", ref="K1:L3"))
    from openpyxl.workbook.defined_name import DefinedName
    try:
        wb.defined_names.add(DefinedName("LOOP", attr_text="LOOP"))  # nome que aponta para si
    except AttributeError:
        wb.defined_names.append(DefinedName("LOOP", attr_text="LOOP"))
    ws["C1"] = "=SUM(A2:A3)+SUM(A5:A12)+T2!D1+T2!F1+T2!G1+T2!I1+T2!L2+LOOP"
    buf = io.BytesIO()
    wb.save(buf)
    texto, n = converter(buf, "C1")
    assert "A2:A12 [Valor] =B2*2" in texto and "  except A4 not in chain" in texto, texto
    assert "D1 =E1" in texto and "E1 =D1" in texto and "H1 array H1:H3" in texto
    assert "- 'T2': unresolved reference Nope!A1 in F1" in texto and "unresolved reference LOOP" in texto and "OFFSET/INDIRECT" in texto
    assert "K2 7" in texto and "L3 =" not in texto and n["falhas"] == 0
    # nome curto nao e coluna; nome de aba ganha do global; constante nao avisa; B1# avisa; meio de matriz
    wb = Workbook()
    s1, s2 = wb.active, wb.create_sheet("S2")
    s1.title = "S1"
    s1["Z1"], s2["A1"], s2["B1"] = 5, 1, 100
    for nome, texto in (("TIR", "S1!$Z$1"), ("Rate", "S2!$B$1"), ("Taxa", "0.05")):
        try:
            wb.defined_names.add(DefinedName(nome, attr_text=texto))
        except AttributeError:
            wb.defined_names.append(DefinedName(nome, attr_text=texto))
    try:
        s1.defined_names.add(DefinedName("Rate", attr_text="S2!$A$1"))
    except AttributeError:
        s1.defined_names["Rate"] = DefinedName("Rate", attr_text="S2!$A$1")
    s1["A1"] = "=TIR*Rate*Taxa+C1"
    s1["C1"] = "=SUM(D1#)"
    s1["E1"] = ArrayFormula("E1:E2", "=Z1:Z2*2")
    buf = io.BytesIO()
    wb.save(buf)
    texto, n = converter(buf, "S1!A1")
    assert "Z1 5" in texto and "A1 1" in texto and "B1 100" not in texto, texto
    assert "unresolved" not in texto and "formula not parsed" in texto and "see Warnings" in texto
    texto, n = converter(buf, "S1!E2")
    assert "E1 array E1:E2" in texto
    import subprocess
    uso = subprocess.run([sys.executable, os.path.abspath(__file__), caminho, "--celula", "Fluxo!D5"],
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    assert uso.returncode == 0 and b"Cadeia: 3 de 618 formulas" in uso.stdout, uso.stdout
    uso = subprocess.run([sys.executable, os.path.abspath(__file__), caminho, "x", "y"],
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    assert uso.returncode == 2 and b"uso:" in uso.stderr
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        assert main(caminho, "'Fluxo'!D5") == 0
        assert os.path.isfile(os.path.join(pasta, "fix.Fluxo_D5.ai.txt"))
        assert main(caminho, "Nada!A1") == 2

    # --diff: duas versoes
    from openpyxl.workbook.defined_name import DefinedName as DN

    def livro(edita=None):
        w = Workbook()
        a = w.active
        a.title = "S"
        a["A1"], a["B1"] = "Item", "Valor"
        for r in range(2, 12):
            a["A%d" % r] = "x%d" % r
            a["B%d" % r] = r
            a["C%d" % r] = "=B%d*2" % r
        a["D1"] = "=SUM(C2:C11)"
        w.create_sheet("T")["A1"] = "=S!D1"
        w.defined_names["Taxa"] = DN("Taxa", attr_text="S!$B$2")
        if edita:
            edita(w)
        b = io.BytesIO()
        w.save(b)
        return b

    velho = livro()
    texto, n = converter_diff(velho, livro())
    assert "No changes found." in texto and n["formulas"] == n["antes"] == n["constantes"] == 0 and n["falhas"] == 0, texto

    def v2(w):
        a = w["S"]
        for r in range(5, 9):
            a["C%d" % r] = "=B%d*3" % r  # bloco de 4 mudou
        a["C11"] = 7  # formula virou constante
        a["B2"] = 99  # constante mudou
        a["E1"] = "novo"  # constante nova
        a["A3"] = None  # constante removida
        w["T"]["A1"] = "=S!D1+1"
        w.create_sheet("U")["A1"] = "=1+1"
        del w["S"].parent.defined_names["Taxa"]
        w.defined_names["Outra"] = DN("Outra", attr_text="S!$B$3")

    texto, n = converter_diff(velho, livro(v2))
    assert n["falhas"] == 0, texto
    assert "C5:C8 =B5*3" in texto and "C5:C8 =B5*2" in texto, texto  # novo e antigo
    assert "C11 =B11*2" in texto and "C11 7" not in texto.split("## Constants")[0]
    assert "B2 2 -> 99" in texto and "E1 added novo" in texto and "A3 removed (was x3)" in texto and "C11 added 7" in texto, texto
    assert "A1 =S!D1+1" in texto and "A1 =S!D1\n" in texto
    assert "- sheet 'U' added (1 formulas)" in texto
    assert "removed: Taxa = S!$B$2" in texto and "added: Outra = S!$B$3" in texto, texto
    assert "C2:C4" not in texto and "SUM" not in texto  # o que nao mudou fica de fora
    assert n["formulas"] == 5 and n["antes"] == 6 and n["constantes"] == 4, n
    # preencher para baixo igual nao e mudanca (mesma formula relativa)
    texto, n = converter_diff(velho, livro(lambda w: w["S"].__setitem__("C2", "=B2*2")))
    assert "No changes found." in texto
    # ruido: "" vira vazio, float em 1e-12 nao conta
    def ruido(w):
        w["S"]["F1"] = ""
        w["S"]["B3"] = 3.0000000000001
    texto, n = converter_diff(velho, livro(ruido))
    assert "No changes found." in texto, texto
    # mais de 200 constantes: corta
    def muitas(w):
        for r in range(1, 301):
            w["T"]["B%d" % r] = r
    texto, n = converter_diff(velho, livro(muitas))
    assert "... 100 more" in texto and texto.count("added") == 200 and n["constantes"] == 300
    # conferencia que falha cai para uma linha por celula
    guarda = confere
    globals()["confere"] = lambda linhas, meu, base=None: False
    try:
        texto, n = converter_diff(velho, livro(v2))
    finally:
        globals()["confere"] = guarda
    assert n["falhas"] > 0 and "WARNING: compact view failed" in texto and "C5 =B5*3" in texto, texto
    # CLI
    pa, pn = os.path.join(pasta, "old v1.xlsx"), os.path.join(pasta, "new v2.xlsx")
    for caminho_x, b in ((pa, velho), (pn, livro(v2))):
        with open(caminho_x, "wb") as arq:
            arq.write(b.getvalue())
    uso = subprocess.run([sys.executable, os.path.abspath(__file__), pa, pn, "--diff"],
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    assert uso.returncode == 0 and b"100%" in uso.stdout, uso.stdout + uso.stderr
    assert os.path.isfile(os.path.join(pasta, "new_v2.vs_old_v1.diff.ai.txt"))
    uso = subprocess.run([sys.executable, os.path.abspath(__file__), pa, os.path.join(pasta, "nada.xlsx"), "--diff"],
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    assert uso.returncode == 2
    uso = subprocess.run([sys.executable, os.path.abspath(__file__), pa, pn, "--celula"],
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    assert uso.returncode == 2 and b"uso:" in uso.stderr

    # pasta inteira
    sub = os.path.join(pasta, "lote")
    os.makedirs(os.path.join(sub, "fundo"))
    for nome_x in ("b.xlsx", "a.XLSM", os.path.join("fundo", "c.xlsx")):
        with open(os.path.join(sub, nome_x), "wb") as arq:
            arq.write(_fixture().getvalue())
    for nome_x in ("~$a.xlsx", "ruim.xlsx", "nota.txt"):
        with open(os.path.join(sub, nome_x), "wb") as arq:
            arq.write(b"nao e planilha")
    out = io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(io.StringIO()):
        assert main_pasta(sub) == 2
    assert os.path.isfile(os.path.join(sub, "a.ai.txt")) and os.path.isfile(os.path.join(sub, "b.ai.txt"))
    assert not os.path.exists(os.path.join(sub, "ruim.ai.txt")) and not os.path.exists(os.path.join(sub, "fundo", "c.ai.txt"))
    assert "Pasta: 3 planilha(s), 1 com problema" in out.getvalue() and "== a.XLSM ==" in out.getvalue(), out.getvalue()
    vazia = os.path.join(pasta, "vazia")
    os.makedirs(vazia)
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        assert main_pasta(vazia) == 2
    uso = subprocess.run([sys.executable, os.path.abspath(__file__), os.path.join(sub, "fundo")],
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    assert uso.returncode == 0 and b"Pasta: 1 planilha(s), 0 com problema" in uso.stdout, uso.stdout

    # macro: zip com vbaProject.bin
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as arq:
        arq.writestr("xl/vbaProject.bin", b"x")
    assert tem_vba(z) and not tem_vba(_fixture())
    # skill do Kiro: instala, pula se ja existe, --forcar sobrescreve
    destino = os.path.join(pasta, "skill")
    with contextlib.redirect_stdout(io.StringIO()):
        assert instalar_skill(destino) == 0
    for rel in ("SKILL.md", "references/format.md", "scripts/excel_ai.py"):
        assert os.path.isfile(os.path.join(destino, rel)), rel
    with open(os.path.join(destino, "SKILL.md"), "w") as f:
        f.write("meu")
    with contextlib.redirect_stdout(io.StringIO()):
        instalar_skill(destino)
    assert open(os.path.join(destino, "SKILL.md")).read() == "meu"
    with contextlib.redirect_stdout(io.StringIO()):
        instalar_skill(destino, forcar=True)
    assert open(os.path.join(destino, "SKILL.md")).read() == SKILL_MD
    print("auto-teste ok")


if __name__ == "__main__":
    if len(sys.argv) == 4 and sys.argv[2] == "--celula":
        sys.exit(main(sys.argv[1], sys.argv[3]))
    if len(sys.argv) == 4 and sys.argv[3] == "--diff":
        sys.exit(main_diff(sys.argv[1], sys.argv[2]))
    if sys.argv[1:2] == ["--instalar-skill"] and len(sys.argv) <= 3:
        sys.exit(instalar_skill(forcar=sys.argv[2:] == ["--forcar"]))
    if len(sys.argv) == 2:
        sys.exit(main_pasta(sys.argv[1]) if os.path.isdir(sys.argv[1]) else main(sys.argv[1]))
    if len(sys.argv) > 1:
        sys.stderr.write("uso: python excel_ai.py planilha.xlsx [--celula Aba!B5]  |  python excel_ai.py antiga.xlsx nova.xlsx --diff  |  python excel_ai.py pasta  |  python excel_ai.py --instalar-skill [--forcar]\n")
        sys.exit(2)
    autoteste()
