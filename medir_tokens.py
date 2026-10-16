"""Mede quantos tokens uma planilha gasta: dump celula a celula vs formulas agrupadas.

Uso:   python medir_tokens.py arquivo.xlsx
Teste: python medir_tokens.py          (sem argumento roda o auto-teste)

Roda 100% local. Imprime so numeros, nenhum conteudo da planilha.
"""
import re
import sys
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import column_index_from_string, get_column_letter

REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def tokens(texto):
    # ponytail: ~4 caracteres por token (estimativa). Use tiktoken se precisar de precisao.
    return len(texto) // 4


def relativa(formula, row, col):
    """=B2*C2 em E2 vira =R[0]C[-3]*R[0]C[-2]: mesmo texto para a coluna inteira."""
    def troca(m):
        cabs, letra, rabs, num = m.groups()
        c, r = column_index_from_string(letra), int(num)
        rs = f"R{r}" if rabs else f"R[{r - row}]"
        cs = f"C{c}" if cabs else f"C[{c - col}]"
        return rs + cs

    try:
        itens = Tokenizer(formula).items
    except Exception:
        return formula
    partes = []
    for t in itens:
        v = t.value
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            aba, sep, ref = v.rpartition("!")
            v = aba + sep + REF.sub(troca, ref)
        partes.append(v)
    return "=" + "".join(partes)


def texto_formula(v):
    """Retorna o texto da formula, ou None se for constante."""
    v = getattr(v, "text", v)  # ArrayFormula (matriz/spill) guarda o texto em .text
    if isinstance(v, str) and v.startswith("="):
        return v
    return None


def medir_aba(ws_f, ws_v):
    valores = {(c.row, c.column): c.value
               for linha in ws_v.iter_rows() for c in linha if c.value is not None}
    bruto, consts = [], []
    grupos = defaultdict(list)  # (coluna, padrao) -> linhas
    n_formulas = 0
    for linha in ws_f.iter_rows():
        for c in linha:
            if c.value is None:
                continue
            ref = f"{get_column_letter(c.column)}{c.row}"
            f = texto_formula(c.value)
            if f is None:
                bruto.append(f"{ref}\t{c.value}")
                consts.append(f"{ref}\t{c.value}")
                continue
            n_formulas += 1
            bruto.append(f"{ref}\t{f}\t{valores.get((c.row, c.column))}")
            grupos[(c.column, relativa(f, c.row, c.column))].append(c.row)

    # Cada sequencia continua de linhas com o mesmo padrao vira uma linha de texto.
    # ponytail: agrupa so na vertical; agrupar em 2D (matriz PDD) se compensar.
    agrupado = []
    for (col, padrao), linhas in grupos.items():
        linhas.sort()
        ini = ant = linhas[0]
        for r in linhas[1:] + [None]:
            if r is not None and r == ant + 1:
                ant = r
                continue
            letra = get_column_letter(col)
            faixa = f"{letra}{ini}" if ini == ant else f"{letra}{ini}:{letra}{ant}"
            agrupado.append(f"{faixa}\t{padrao}")
            if r is not None:
                ini = ant = r

    return {
        "formulas": n_formulas,
        "blocos": len(agrupado),
        "bruto": tokens("\n".join(bruto)),
        "agrupado": tokens("\n".join(agrupado)),
        "consts": tokens("\n".join(consts)),
    }


def main(caminho):
    wb_f = load_workbook(caminho, read_only=True)
    wb_v = load_workbook(caminho, read_only=True, data_only=True)
    print(f"{'aba':<20}{'formulas':>9}{'blocos':>8}{'bruto':>10}{'agrupado':>10}{'constantes':>12}")
    tot = defaultdict(int)
    for nome in wb_f.sheetnames:
        r = medir_aba(wb_f[nome], wb_v[nome])
        for k, v in r.items():
            tot[k] += v
        print(f"{nome[:19]:<20}{r['formulas']:>9}{r['blocos']:>8}{r['bruto']:>10}{r['agrupado']:>10}{r['consts']:>12}")

    novo = tot["agrupado"] + tot["consts"]
    print()
    print(f"Tokens hoje (celula a celula):      {tot['bruto']:>10,}")
    print(f"Tokens com formulas agrupadas:      {novo:>10,}")
    print(f"  so a parte das formulas:          {tot['agrupado']:>10,}")
    print(f"  constantes (dados, nao comprime): {tot['consts']:>10,}")
    if novo:
        print(f"Reducao total: {tot['bruto'] / novo:.1f}x")
    form_bruto = tot["bruto"] - tot["consts"]
    if tot["agrupado"]:
        print(f"Reducao so nas formulas: {form_bruto / tot['agrupado']:.1f}x")


def autoteste():
    assert relativa("=B2*C2", 2, 5) == "=R[0]C[-3]*R[0]C[-2]"
    assert relativa("=B3*C3", 3, 5) == relativa("=B2*C2", 2, 5)
    assert relativa("=B2*$A$1", 2, 3) == "=R[0]C[-1]*R1C1"
    assert relativa("=Fluxo!B2", 2, 2) == "=Fluxo!R[0]C[0]"
    assert relativa("='Plan 1'!B2", 2, 2) == "='Plan 1'!R[0]C[0]"
    # xlsx guarda formula em ingles com virgula; texto entre aspas nao muda
    assert relativa('=IF(A2="B2",1,0)', 2, 2) == '=IF(R[0]C[-1]="B2",1,0)'
    assert relativa("=_xlfn.XLOOKUP(A2,Base!A:A,Base!C:C)", 2, 2) == "=_xlfn.XLOOKUP(R[0]C[-1],Base!A:A,Base!C:C)"
    print("auto-teste ok")


if __name__ == "__main__":
    main(sys.argv[1]) if len(sys.argv) > 1 else autoteste()
